from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

ROOT=Path('/mnt/data/MTI-CardioOncology')
out=ROOT/'manuscript'/'JCF_MTI_Brief_Report_BLINDED.docx'
md=(ROOT/'manuscript'/'JCF_MTI_Brief_Report.md').read_text()

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def add_text(p, text, bold=False, italic=False, size=None):
    # Basic markdown emphasis parser.
    parts=re.split(r'(\*\*.*?\*\*|\*.*?\*)',text)
    for part in parts:
        if not part: continue
        b=bold; i=italic; val=part
        if part.startswith('**') and part.endswith('**'):
            val=part[2:-2]; b=True
        elif part.startswith('*') and part.endswith('*'):
            val=part[1:-1]; i=True
        r=p.add_run(val); r.bold=b; r.italic=i
        if size: r.font.size=Pt(size)
        r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')

def style_doc(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(1);sec.bottom_margin=Inches(1);sec.left_margin=Inches(1);sec.right_margin=Inches(1)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Times New Roman';normal.font.size=Pt(12); normal._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman');normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    normal.paragraph_format.line_spacing=2; normal.paragraph_format.space_after=Pt(0)
    for s in ['Heading 1','Heading 2']:
        st=styles[s];st.font.name='Times New Roman';st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman');st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman');st.font.bold=True;st.font.color.rgb=None
    styles['Heading 1'].font.size=Pt(13); styles['Heading 2'].font.size=Pt(12)
    # page numbers
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('Page '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table=doc.add_table(rows=1, cols=len(headers)); table.style='Table Grid'; table.autofit=False
    hdr=table.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text=''; p=hdr[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_text(p,h,bold=True,size=font_size);set_cell_shading(hdr[i],'E6ECEE');set_cell_margins(hdr[i])
    for row in rows:
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=''; p=cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;add_text(p,str(val),size=font_size);set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return table

doc=Document();style_doc(doc)
# Title
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_text(p,'Mechanical Transition Instability on Pretreatment Tissue Doppler Imaging Identifies Risk of Cancer Therapy-Related Cardiac Dysfunction in Women With Breast Cancer',bold=True,size=15)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_text(p,'Brief Report',italic=True,size=11)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_text(p,'Keywords: cardio-oncology; cardiotoxicity; echocardiography; tissue Doppler imaging',size=10)

# Parse manuscript up to references, skipping markdown title.
lines=md.splitlines(); in_refs=False; refs=[]
for line in lines[1:]:
    if line.startswith('## References'):
        in_refs=True; continue
    if in_refs:
        if line.strip(): refs.append(line.strip())
        continue
    if line.startswith('## '):
        h=line[3:].strip(); doc.add_heading(h,level=1); continue
    if not line.strip(): continue
    p=doc.add_paragraph(); add_text(p,line.strip())

# References
p=doc.add_heading('References',level=1)
for r in refs:
    p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);add_text(p,r)

# Figure legends
leg=(ROOT/'manuscript'/'Figure_Legends.md').read_text().splitlines()
doc.add_page_break();doc.add_heading('Figure Legends',level=1)
for line in leg:
    if not line.strip() or line.startswith('#'): continue
    p=doc.add_paragraph();add_text(p,line)

# Table 1

doc.add_page_break();doc.add_heading('Table 1. Cohort characteristics, MTI phenotype, cross-validated performance, and association with CTRCD',level=1)
add_text(doc.add_paragraph(),'Panel A. Pretreatment TDI analysis cohort',bold=True)
rowsA=[
['Age, y','58 [49-66]','61 [54-64]','57 [49-66]'],['Heart rate, beats/min','72.5 [64.3-80.0]','78 [71-87]','72 [64-79]'],['LVEF, %','65.7 [61.2-70.9]','61.4 [58.4-66.9]','65.8 [61.6-72.0]'],['LV end-diastolic diameter, cm','4.29 [3.97-4.63]','4.22 [3.81-4.64]','4.30 [3.97-4.63]'],['LV end-systolic diameter, cm','2.70 [2.44-2.94]','2.85 [2.37-3.12]','2.69 [2.45-2.93]'],['Follow-up, d','437 [220-874]','316 [209-720]','453 [224-879]'],['Anthracycline therapy','163/248 (65.7%)','20/24 (83.3%)','143/224 (63.8%)'],['Anti-HER2 therapy','158/248 (63.7%)','17/24 (70.8%)','141/224 (62.9%)'],['Hypertension','62/247 (25.1%)','6/24 (25.0%)','56/223 (25.1%)'],['Dyslipidemia','53/248 (21.4%)','6/24 (25.0%)','47/224 (21.0%)'],['Diabetes mellitus','14/248 (5.6%)','0/24 (0%)','14/224 (6.2%)'],['Prior anthracycline exposure','29/248 (11.7%)','7/24 (29.2%)','22/224 (9.8%)'],['Prior anti-HER2 exposure','7/248 (2.8%)','3/24 (12.5%)','4/224 (1.8%)']]
add_table(doc,['Characteristic','Overall (n=270)','CTRCD (n=27)','No CTRCD (n=243)'],rowsA,[2.6,1.3,1.3,1.4],8.2)
p=doc.add_paragraph();add_text(p,'Values are median [IQR] or n/N (%). Binary treatment/comorbidity variables contain missing values in the public dataset; denominators are shown explicitly.',size=8)

add_text(doc.add_paragraph(),'Panel B. Four-component Mechanical Transition Instability phenotype',bold=True)
rowsB=[
['Transition slope deficit','Reduced steepness from peak systolic velocity toward zero','120.5 [103.3-164.1]','151.0 [122.4-181.9]','0.650 (0.528-0.773)','0.011'],
['Transition curvature deficit','Blunting of the systolic-to-relaxation bend','742.0 [540.3-930.9]','909.9 [717.7-1146.5]','0.652 (0.531-0.766)','0.010'],
['Early-relaxation latency','Delay from systolic termination to E′ peak','0.175 [0.150-0.200]','0.152 [0.132-0.175]','0.632 (0.510-0.742)','0.024'],
['E′ attenuation','Reduced absolute E′ peak velocity','9.15 [7.89-10.48]','10.51 [8.11-12.59]','0.621 (0.519-0.714)','0.039']]
add_table(doc,['MTI component','Operational meaning','CTRCD median [IQR]','No CTRCD median [IQR]','Event-oriented AUROC (95% CI)','P'],rowsB,[1.25,1.9,1.15,1.2,1.25,.55],7.6)
p=doc.add_paragraph();add_text(p,'P values are exploratory two-sided Mann-Whitney tests. Higher oriented component values indicate greater abnormality when constructing MTI.',size=8)

add_text(doc.add_paragraph(),'Panel C. Repeated cross-fitted discrimination',bold=True)
rowsC=[['MTI','0.695 (0.590-0.794)','0.192 (0.141-0.321)','0.677 (0.570-0.784)','0.657'],['Baseline LVEF','0.657 (0.549-0.759)','0.166 (0.121-0.281)','0.678 (0.578-0.767)','0.672'],['Published E′ timing/velocity comparator','0.663 (0.574-0.747)','0.150 (0.120-0.241)','0.660 (0.570-0.750)','0.637'],['MTI + LVEF','0.712 (0.608-0.807)','0.211 (0.153-0.357)','0.715 (0.611-0.810)','0.707']]
add_table(doc,['Model','AUROC (95% CI)','AUPRC (95% CI)','Harrell C-index (95% CI)','1-y AUROC'],rowsC,[2.0,1.4,1.4,1.45,.85],8)
p=doc.add_paragraph();add_text(p,'The paired difference for MTI + LVEF versus LVEF was +0.055 for AUROC (95% CI, -0.040 to 0.153) and +0.037 for C-index (95% CI, -0.059 to 0.137); superiority was not established.',size=8)

add_text(doc.add_paragraph(),'Panel D. Association with time to CTRCD and robustness',bold=True)
rowsD=[['MTI unadjusted','270 / 27','2.37 (1.49-3.78)','<0.001'],['Adjusted for LVEF','270 / 27','2.20 (1.36-3.55)','0.001'],['Adjusted for age + LVEF','270 / 27','2.19 (1.34-3.58)','0.002'],['Adjusted for age + LVEF + anthracycline + anti-HER2','248 / 24','1.95 (1.19-3.18)','0.008'],['Baseline LVEF >=50%, adjusted for LVEF','261 / 26','2.22 (1.41-3.51)','<0.001']]
add_table(doc,['Analysis','n / events','MTI HR per 1 SD (95% CI)','P'],rowsD,[3.4,1.0,1.65,.7],8.2)
foot='CTRCD occurred in 5/90, 6/90, and 16/90 women in increasing MTI tertiles. A restricted-cubic-spline model showed no evidence of departure from a linear MTI-risk association (P for nonlinearity=0.999). In an exploratory surveillance-efficiency analysis using held-out MTI + LVEF risk rankings, the highest-risk 20% contained 13/27 future events (48.1%) and had PPV 24.1%, compared with 10.0% cohort prevalence. Across Savitzky-Golay windows of 21 to 61 samples, MTI AUROC ranged from 0.686 to 0.710 and C-index from 0.673 to 0.694. Women with complete TDI differed from women without processed TDI in age (SMD 0.45), follow-up duration (SMD -0.41), anti-HER2 exposure (SMD -0.34), and hypertension (SMD 0.22), while CTRCD prevalence was similar (SMD -0.01). Abbreviations: AUPRC, area under the precision-recall curve; AUROC, area under the receiver-operating-characteristic curve; CTRCD, cancer therapy-related cardiac dysfunction; HR, hazard ratio; LVEF, left ventricular ejection fraction; MTI, Mechanical Transition Instability; TDI, Tissue Doppler imaging.'
p=doc.add_paragraph();add_text(p,foot,size=8)

doc.core_properties.title='Mechanical Transition Instability on Pretreatment Tissue Doppler Imaging and Risk of CTRCD'
doc.core_properties.subject='JCF Brief Report blinded manuscript'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.save(out)
print(out)
